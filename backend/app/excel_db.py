import os
import shutil
import threading
from datetime import datetime
from typing import List, Optional, Dict, Any
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import Property, PropertyCreate, PropertyUpdate, Lead, LeadCreate, LeadUpdate

DEFAULT_DB_PATH = Path(__file__).resolve().parent.parent.parent / "data" / "real_estate_database.xlsx"

PROPERTY_COLUMNS = [
    "Property_ID", "Date_Added", "Listing_Type", "BHK", "Building_Name",
    "Location_Area", "Furnishing", "Rent_Price", "Deposit", "Features",
    "Contact_Numbers", "Agency", "Status", "Source", "Raw_Message"
]

LEAD_COLUMNS = [
    "Lead_ID", "Date_Added", "Name", "Phone", "Client_Type",
    "Preferred_BHK", "Preferred_Locations", "Min_Budget", "Max_Budget",
    "Furnishing", "Move_In_Date", "Status", "Notes"
]

class ExcelDatabase:
    def __init__(self, file_path: Optional[Path] = None):
        self.file_path = file_path or DEFAULT_DB_PATH
        self.lock = threading.Lock()
        self._ensure_workbook()

    def _ensure_workbook(self):
        """Creates the Excel workbook with proper sheets and styled headers if missing."""
        self.file_path.parent.mkdir(parents=True, exist_ok=True)
        if not self.file_path.exists():
            wb = openpyxl.Workbook()
            # Setup Properties sheet
            ws_props = wb.active
            ws_props.title = "Properties"
            self._apply_header_style(ws_props, PROPERTY_COLUMNS, header_color="1E3A8A")

            # Setup Clients_CRM sheet
            ws_leads = wb.create_sheet(title="Clients_CRM")
            self._apply_header_style(ws_leads, LEAD_COLUMNS, header_color="065F46")

            wb.save(self.file_path)

    def _apply_header_style(self, ws, columns: List[str], header_color: str):
        ws.append(columns)
        font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
        alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )
        ws.row_dimensions[1].height = 28
        for col_idx, col_name in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = font
            cell.fill = fill
            cell.alignment = alignment
            cell.border = border
            ws.column_dimensions[get_column_letter(col_idx)].width = max(len(col_name) + 6, 14)

    def _safe_save(self, wb):
        try:
            wb.save(self.file_path)
        except PermissionError:
            # If Excel has the file open and locked, save to a backup copy
            backup_path = self.file_path.parent / f"{self.file_path.stem}_backup.xlsx"
            wb.save(backup_path)
            raise RuntimeError(
                f"Notice: '{self.file_path.name}' is currently locked by another application (e.g. MS Excel). "
                f"Changes saved to '{backup_path.name}'. Please close the file in Excel to allow direct sync."
            )

    # ---------------- Properties Operations ---------------- #

    def get_all_properties(self) -> List[Property]:
        with self.lock:
            if not self.file_path.exists():
                return []
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            if "Properties" not in wb.sheetnames:
                return []
            ws = wb["Properties"]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) <= 1:
                return []
            
            headers = [str(h) for h in rows[0]]
            properties = []
            for r in rows[1:]:
                if not any(r):
                    continue
                row_dict = dict(zip(headers, r))
                prop_id = str(row_dict.get("Property_ID") or "")
                if not prop_id:
                    continue
                properties.append(
                    Property(
                        id=prop_id,
                        date_added=str(row_dict.get("Date_Added") or datetime.now().strftime("%Y-%m-%d")),
                        listing_type=str(row_dict.get("Listing_Type") or "Rent"),
                        bhk=str(row_dict.get("BHK") or "2 BHK"),
                        building_name=str(row_dict.get("Building_Name") or ""),
                        location_area=str(row_dict.get("Location_Area") or ""),
                        furnishing=str(row_dict.get("Furnishing") or "Semi-Furnished"),
                        rent_price=float(row_dict.get("Rent_Price") or 0.0),
                        deposit=str(row_dict.get("Deposit") or ""),
                        features=str(row_dict.get("Features") or ""),
                        contact_numbers=str(row_dict.get("Contact_Numbers") or ""),
                        agency=str(row_dict.get("Agency") or ""),
                        status=str(row_dict.get("Status") or "Available"),
                        source=str(row_dict.get("Source") or "Manual"),
                        raw_message=str(row_dict.get("Raw_Message") or "")
                    )
                )
            return properties

    def get_property(self, prop_id: str) -> Optional[Property]:
        props = self.get_all_properties()
        for p in props:
            if p.id.lower() == prop_id.lower():
                return p
        return None

    def add_property(self, prop: PropertyCreate) -> Property:
        with self.lock:
            wb = openpyxl.load_workbook(self.file_path)
            if "Properties" not in wb.sheetnames:
                ws = wb.create_sheet(title="Properties")
                self._apply_header_style(ws, PROPERTY_COLUMNS, "1E3A8A")
            else:
                ws = wb["Properties"]

            # Compute next ID
            next_num = 1
            for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
                val = row[0]
                if val and str(val).startswith("PROP-"):
                    try:
                        num = int(str(val).split("-")[-1])
                        if num >= next_num:
                            next_num = num + 1
                    except ValueError:
                        pass
            
            new_id = f"PROP-{next_num:04d}"
            today_str = datetime.now().strftime("%Y-%m-%d")

            new_row = [
                new_id,
                today_str,
                prop.listing_type,
                prop.bhk,
                prop.building_name,
                prop.location_area,
                prop.furnishing,
                prop.rent_price,
                prop.deposit,
                prop.features,
                prop.contact_numbers,
                prop.agency,
                prop.status,
                prop.source,
                prop.raw_message or ""
            ]
            ws.append(new_row)
            self._safe_save(wb)

            return Property(
                id=new_id,
                date_added=today_str,
                **prop.model_dump()
            )

    def update_property(self, prop_id: str, update: PropertyUpdate) -> Optional[Property]:
        with self.lock:
            wb = openpyxl.load_workbook(self.file_path)
            if "Properties" not in wb.sheetnames:
                return None
            ws = wb["Properties"]

            headers = [cell.value for cell in ws[1]]
            target_row_idx = None
            for row_idx in range(2, ws.max_row + 1):
                cell_val = str(ws.cell(row=row_idx, column=1).value or "")
                if cell_val.lower() == prop_id.lower():
                    target_row_idx = row_idx
                    break

            if not target_row_idx:
                return None

            update_data = update.model_dump(exclude_unset=True)
            field_col_map = {
                "listing_type": "Listing_Type",
                "bhk": "BHK",
                "building_name": "Building_Name",
                "location_area": "Location_Area",
                "furnishing": "Furnishing",
                "rent_price": "Rent_Price",
                "deposit": "Deposit",
                "features": "Features",
                "contact_numbers": "Contact_Numbers",
                "agency": "Agency",
                "status": "Status"
            }

            for field_key, val in update_data.items():
                col_name = field_col_map.get(field_key)
                if col_name and col_name in headers:
                    col_idx = headers.index(col_name) + 1
                    ws.cell(row=target_row_idx, column=col_idx).value = val

            self._safe_save(wb)
        return self.get_property(prop_id)

    def delete_property(self, prop_id: str) -> bool:
        with self.lock:
            wb = openpyxl.load_workbook(self.file_path)
            if "Properties" not in wb.sheetnames:
                return False
            ws = wb["Properties"]
            target_row_idx = None
            for row_idx in range(2, ws.max_row + 1):
                cell_val = str(ws.cell(row=row_idx, column=1).value or "")
                if cell_val.lower() == prop_id.lower():
                    target_row_idx = row_idx
                    break

            if target_row_idx:
                ws.delete_rows(target_row_idx, 1)
                self._safe_save(wb)
                return True
            return False

    # ---------------- Leads / CRM Operations ---------------- #

    def get_all_leads(self) -> List[Lead]:
        with self.lock:
            if not self.file_path.exists():
                return []
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
            if "Clients_CRM" not in wb.sheetnames:
                return []
            ws = wb["Clients_CRM"]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) <= 1:
                return []

            headers = [str(h) for h in rows[0]]
            leads = []
            for r in rows[1:]:
                if not any(r):
                    continue
                row_dict = dict(zip(headers, r))
                lead_id = str(row_dict.get("Lead_ID") or "")
                if not lead_id:
                    continue
                leads.append(
                    Lead(
                        id=lead_id,
                        date_added=str(row_dict.get("Date_Added") or datetime.now().strftime("%Y-%m-%d")),
                        name=str(row_dict.get("Name") or ""),
                        phone=str(row_dict.get("Phone") or ""),
                        client_type=str(row_dict.get("Client_Type") or "Tenant"),
                        preferred_bhk=str(row_dict.get("Preferred_BHK") or "2 BHK"),
                        preferred_locations=str(row_dict.get("Preferred_Locations") or ""),
                        min_budget=float(row_dict.get("Min_Budget") or 0.0),
                        max_budget=float(row_dict.get("Max_Budget") or 0.0),
                        furnishing=str(row_dict.get("Furnishing") or "Any"),
                        move_in_date=str(row_dict.get("Move_In_Date") or "Immediate"),
                        status=str(row_dict.get("Status") or "New"),
                        notes=str(row_dict.get("Notes") or "")
                    )
                )
            return leads

    def get_lead(self, lead_id: str) -> Optional[Lead]:
        leads = self.get_all_leads()
        for l in leads:
            if l.id.lower() == lead_id.lower():
                return l
        return None

    def add_lead(self, lead: LeadCreate) -> Lead:
        with self.lock:
            wb = openpyxl.load_workbook(self.file_path)
            if "Clients_CRM" not in wb.sheetnames:
                ws = wb.create_sheet(title="Clients_CRM")
                self._apply_header_style(ws, LEAD_COLUMNS, "065F46")
            else:
                ws = wb["Clients_CRM"]

            next_num = 1
            for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
                val = row[0]
                if val and str(val).startswith("LEAD-"):
                    try:
                        num = int(str(val).split("-")[-1])
                        if num >= next_num:
                            next_num = num + 1
                    except ValueError:
                        pass

            new_id = f"LEAD-{next_num:04d}"
            today_str = datetime.now().strftime("%Y-%m-%d")

            new_row = [
                new_id,
                today_str,
                lead.name,
                lead.phone,
                lead.client_type,
                lead.preferred_bhk,
                lead.preferred_locations,
                lead.min_budget,
                lead.max_budget,
                lead.furnishing,
                lead.move_in_date,
                lead.status,
                lead.notes
            ]
            ws.append(new_row)
            self._safe_save(wb)

            return Lead(
                id=new_id,
                date_added=today_str,
                **lead.model_dump()
            )

    def update_lead(self, lead_id: str, update: LeadUpdate) -> Optional[Lead]:
        with self.lock:
            wb = openpyxl.load_workbook(self.file_path)
            if "Clients_CRM" not in wb.sheetnames:
                return None
            ws = wb["Clients_CRM"]

            headers = [cell.value for cell in ws[1]]
            target_row_idx = None
            for row_idx in range(2, ws.max_row + 1):
                cell_val = str(ws.cell(row=row_idx, column=1).value or "")
                if cell_val.lower() == lead_id.lower():
                    target_row_idx = row_idx
                    break

            if not target_row_idx:
                return None

            update_data = update.model_dump(exclude_unset=True)
            field_col_map = {
                "name": "Name",
                "phone": "Phone",
                "client_type": "Client_Type",
                "preferred_bhk": "Preferred_BHK",
                "preferred_locations": "Preferred_Locations",
                "min_budget": "Min_Budget",
                "max_budget": "Max_Budget",
                "furnishing": "Furnishing",
                "move_in_date": "Move_In_Date",
                "status": "Status",
                "notes": "Notes"
            }

            for field_key, val in update_data.items():
                col_name = field_col_map.get(field_key)
                if col_name and col_name in headers:
                    col_idx = headers.index(col_name) + 1
                    ws.cell(row=target_row_idx, column=col_idx).value = val

            self._safe_save(wb)
        return self.get_lead(lead_id)

    def delete_lead(self, lead_id: str) -> bool:
        with self.lock:
            wb = openpyxl.load_workbook(self.file_path)
            if "Clients_CRM" not in wb.sheetnames:
                return False
            ws = wb["Clients_CRM"]
            target_row_idx = None
            for row_idx in range(2, ws.max_row + 1):
                cell_val = str(ws.cell(row=row_idx, column=1).value or "")
                if cell_val.lower() == lead_id.lower():
                    target_row_idx = row_idx
                    break

            if target_row_idx:
                ws.delete_rows(target_row_idx, 1)
                self._safe_save(wb)
                return True
            return False

# Global instance
excel_db = ExcelDatabase()
